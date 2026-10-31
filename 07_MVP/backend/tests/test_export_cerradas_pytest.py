"""RF-30 / RF-31 (T-027): Excel de resultados de preguntas cerradas. mongomock, sin IA.

Se abre el archivo con openpyxl y se compara con `closed-results` (mismos parámetros): una sola fuente de cifras.
"""
from __future__ import annotations

import io
import json
import os
import sys
from pathlib import Path
from urllib.parse import quote

import mongomock
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ.setdefault("ANTHROPIC_API_KEY", "sk-test")

from app import database  # noqa: E402

database._client = mongomock.MongoClient()

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.chart import BarChart, LineChart  # noqa: E402

from app import store  # noqa: E402
from app.auth import create_token, hash_password  # noqa: E402
from app.database import get_db  # noqa: E402
from app.main import app  # noqa: E402
from app.models import User, user_to_doc  # noqa: E402
from app.services import export_closed, quantitative  # noqa: E402

N = 40
DF = pd.DataFrame({
    "Satisfacción": [str(i % 5 + 1) for i in range(N)],
    "Recomendación": [str((i * 3) % 11) for i in range(N)],
    "Recomienda": ["Sí" if i % 3 else "No" for i in range(N)],
    "Color": [["Rojo", "Verde", "Azul"][i % 3] for i in range(N)],
    "Frutas": ["Manzana, Pera" if i % 2 else "Pera" for i in range(N)],
    "Edad": [str(18 + i) for i in range(N)],
    "Género": [["Mujer", "Hombre", "Otro"][i % 3] for i in range(N)],
    "Tramo": ["20-24" if i < 5 else "25-29" if i < 12 else "30-34" if i < 25 else "35-39" for i in range(N)],
    "Comuna": [f"C{i % 20:02d}" for i in range(N)],
    "Opción": [f"Opc {i % 12}" for i in range(N)],
})
TIPOS = {"Satisfacción": "escala", "Recomendación": "nps", "Recomienda": "si_no", "Color": "seleccion_unica",
         "Frutas": "seleccion_multiple", "Edad": "numerica", "Género": "segmento", "Tramo": "segmento",
         "Comuna": "seleccion_unica", "Opción": "seleccion_unica"}
PREGUNTAS = [{"titulo": t, "columnas": [t], "tipo": tp} for t, tp in TIPOS.items()]


@pytest.fixture(scope="module")
def env():
    c = TestClient(app, raise_server_exceptions=False)
    heads, orgs = {}, {}
    for who in ("a", "b"):
        r = c.post("/auth/signup", json={"email": f"xcer-{who}@x.cl", "password": "clave12345"})
        heads[who] = {"Authorization": f"Bearer {r.json()['token']}"}
        orgs[who] = r.json()["user"]["organizacion"]["id"]
    u = User(id="usr_lectorxcer", email="xcer-lector@x.cl", password_hash=hash_password("clave12345"), nombre="L",
             empresa="", created_at=store.now_iso(), org_id=orgs["a"], rol="lector")
    get_db().users.insert_one(user_to_doc(u))
    r = c.post("/api/projects", headers=heads["a"], data={"name": "Encuesta Ñandú", "preguntas": json.dumps(PREGUNTAS)},
               files={"file": ("e.csv", io.BytesIO(DF.to_csv(index=False).encode()), "text/csv")})
    assert r.status_code == 200, r.text
    return {"c": c, "h": heads["a"], "hb": heads["b"], "hl": {"Authorization": f"Bearer {create_token(u.id)}"},
            "pid": r.json()["id"], "ids": {p["titulo"]: p["id"] for p in c.get(
                f"/api/projects/{r.json()['id']}/closed-results", headers=heads["a"]).json()["preguntas"]}}


def _query(**kw) -> str:
    return "&".join(f"{k}={quote(v)}" for k, v in kw.items() if v is not None)


def _xlsx(env, por=None, cortes=None, unir=None, h="h"):
    q = _query(por=env["ids"].get(por, por) if por else None, cortes=cortes, unir=unir)
    r = env["c"].get(f"/api/projects/{env['pid']}/export/closed-xlsx?{q}", headers=env[h])
    return r, (load_workbook(io.BytesIO(r.content)) if r.status_code == 200 else None)


def _json(env, por=None, cortes=None, unir=None):
    q = _query(por=env["ids"].get(por, por) if por else None, cortes=cortes, unir=unir)
    return env["c"].get(f"/api/projects/{env['pid']}/closed-results?{q}", headers=env["h"]).json()


def _bloques(ws):
    """título -> fila donde empieza el bloque (las celdas de la columna A en negrita y tamaño 12)."""
    return {c.value: c.row for c in ws["A"] if c.font and c.font.size == 12 and c.font.bold}


def _titulo_grafico(ch):
    return ch.title.tx.rich.p[0].r[0].t


def _graficos(ws):
    return {_titulo_grafico(ch): ch for ch in ws._charts}


def _tabla(ws, fila, col=1, ancho=3):
    """Filas desde `fila` hasta la primera vacía (en la columna `col`), con `ancho` columnas."""
    out = []
    while ws.cell(fila, col).value is not None:
        out.append([ws.cell(fila, col + k).value for k in range(ancho)])
        fila += 1
    return out


# ---------- cifras idénticas a closed-results ----------

def test_totales_sin_por_coinciden_con_closed_results_y_hojas(env):
    r, wb = _xlsx(env)
    assert r.status_code == 200 and wb.sheetnames == ["Resumen", "Totales"]
    res = _json(env)
    assert len(res["preguntas"]) == 10
    ws, filas = wb["Totales"], _bloques(wb["Totales"])
    for p in res["preguntas"]:
        i = filas[p["titulo"]]
        assert [ws.cell(i + k, 2).value for k in (1, 2, 3, 4)] == [export_closed.TIPOS[p["tipo"]], p["n_respuestas"], p["n_na"], p["tasa_respuesta"]]
        if p.get("opciones"):
            assert _tabla(ws, i + 7)[:len(p["opciones"])] == [[o["valor"], o["n"], o["pct"]] for o in p["opciones"]]
    # estadísticos por tipo
    def extras(titulo):
        p = next(p for p in res["preguntas"] if p["titulo"] == titulo)
        i, ops = filas[titulo], len(p.get("opciones", []))
        return p, _tabla(ws, i + 6 + ops + 2)
    p, ex = extras("Satisfacción")
    assert ex[0] == ["Promedio", p["escala"]["promedio"], None] and ex[1][1] == p["escala"]["desviacion"]
    assert ex[2] == ["Notas altas (4 y 5)", p["escala"]["top2"]["n"], p["escala"]["top2"]["pct"]]
    assert ex[3][1:] == [p["escala"]["bottom2"]["n"], p["escala"]["bottom2"]["pct"]]
    p, ex = extras("Recomendación")
    assert ex[0][1] == p["nps"]["valor"] and ex[1] == ["Promotores", p["nps"]["promotores"]["n"], p["nps"]["promotores"]["pct"]]
    assert ex[3][0] == "Detractores"
    p = next(p for p in res["preguntas"] if p["titulo"] == "Edad")
    i = filas["Edad"]
    assert _tabla(ws, i + 6) == [["Promedio", p["numerica"]["promedio"], None], ["Mediana", p["numerica"]["mediana"], None],
                                 ["Mínimo", 18, None], ["Máximo", 57, None]]
    p, ex = extras("Frutas")
    assert ex == [[quantitative.NOTA_MULTIPLE, None, None]]
    resumen = wb["Resumen"]
    assert resumen["B3"].value == "Encuesta Ñandú" and resumen["B5"].value == N
    assert resumen["A6"].value == export_closed.NOTA


def test_por_grupo_coincide_con_closed_results_con_y_sin_unir_y_con_cortes(env):
    casos = [dict(por="Género"), dict(por="Tramo"),
             dict(por="Tramo", unir=json.dumps([{"nombre": "20-29", "valores": ["20-24", "25-29"]}])),
             dict(por="Edad", cortes="18,30,40")]
    for kw in casos:
        r, wb = _xlsx(env, **kw)
        assert r.status_code == 200 and wb.sheetnames == ["Resumen", "Totales", "Por grupo"], kw
        res = _json(env, **kw)
        ws, filas = wb["Por grupo"], _bloques(wb["Por grupo"])
        assert {p["titulo"] for p in res["preguntas"]} == {t for t in filas}, kw
        for p in res["preguntas"]:
            tabla = _tabla(ws, filas[p["titulo"]] + 2, ancho=40)
            assert [f[0] for f in tabla] == ["Todas las personas"] + [g["valor"] for g in res["grupos"]]
            for fila, item, g in zip(tabla, [p] + p["por_grupo"], [{}] + res["grupos"]):
                assert fila[1] == item["n_respuestas"] and (fila[2] == "Sí") == g.get("pocos_datos", False)
                if p["tipo"] == "escala":
                    assert fila[3:6] == [item["escala"]["promedio"], item["escala"]["top2"]["pct"], item["escala"]["bottom2"]["pct"]]
                elif p["tipo"] == "nps":
                    assert fila[3:6] == [item["nps"]["valor"], item["nps"]["promotores"]["pct"], item["nps"]["detractores"]["pct"]]
                elif p["tipo"] == "numerica":
                    n = item["numerica"]
                    assert fila[3:7] == [n["promedio"], n["mediana"], n["min"], n["max"]]
                else:  # todas las opciones, en el orden del total
                    assert fila[3:3 + len(item["opciones"])] == [o["pct"] for o in item["opciones"]]


# ---------- gráficos ----------

def test_linea_en_escala_por_numerica_con_tramos_y_barras_sin_orden(env):
    _, wb = _xlsx(env, por="Edad", cortes="18,30,40")
    g = _graficos(wb["Por grupo"])
    linea = g["Promedio por grupo (tendencia)"]
    assert isinstance(linea, LineChart) and len(linea.series) == 1 and linea.dataLabels.showVal
    assert any(isinstance(c, LineChart) and _titulo_grafico(c) == "NPS por grupo (tendencia)" for c in wb["Por grupo"]._charts)
    # género no tiene orden: barras (columnas) y sin tendencia
    _, wb = _xlsx(env, por="Género")
    g = _graficos(wb["Por grupo"])
    barras = g["Promedio por grupo"]
    assert isinstance(barras, BarChart) and barras.type == "col" and barras.dataLabels.showVal
    assert not any(isinstance(c, LineChart) for c in wb["Por grupo"]._charts)
    # tramos con número en el nombre (categórica ordenada) sí son tendencia
    _, wb = _xlsx(env, por="Tramo")
    assert "Promedio por grupo (tendencia)" in _graficos(wb["Por grupo"])
    assert quantitative.orden_natural(_json(env, por="Tramo")) and not quantitative.orden_natural(_json(env, por="Género"))


def test_barras_agrupadas_en_categoricas_por_grupo_sin_la_fila_todas(env):
    _, wb = _xlsx(env, por="Género")
    ws = wb["Por grupo"]
    ch = next(c for c in ws._charts if _titulo_grafico(c) == "Porcentaje de cada opción por grupo")
    assert isinstance(ch, BarChart) and ch.grouping == "clustered" and ch.dataLabels.showVal
    assert len(ch.series) <= 6 and len(ch.series) >= 2
    for s in ch.series:  # las barras cubren solo los 3 grupos, nunca "Todas las personas"
        a, b = s.val.numRef.f.split("!")[1].replace("$", "").split(":")
        assert int(b[1:]) - int(a[1:]) + 1 == 3
    assert ws.cell(int(a[1:]) - 1, 1).value == "Todas las personas"
    # una pregunta con más de 6 opciones lo dice en el título
    _, wb = _xlsx(env, por="Género")
    assert any("6 primeras opciones" in t for t in _graficos(wb["Por grupo"]))


def test_mas_de_15_grupos_sin_grafico_y_con_el_texto(env):
    _, wb = _xlsx(env, por="Comuna")
    ws = wb["Por grupo"]
    assert ws._charts == []
    assert sum(1 for c in ws["A"] if c.value == export_closed.DEMASIADOS) == len(_json(env, por="Comuna")["preguntas"])


def test_mas_de_10_opciones_el_grafico_dice_las_10_mas_frecuentes(env):
    _, wb = _xlsx(env)
    g = _graficos(wb["Totales"])
    ch = g["10 opciones más frecuentes de 12"]
    assert isinstance(ch, BarChart) and ch.type == "bar" and ch.dataLabels.showVal
    assert ch.x_axis.scaling.orientation == "maxMin"
    assert "Porcentaje por opción" in g  # las demás, con 10 opciones o menos, sin recorte
    assert "10 opciones más frecuentes de 11" not in g  # el NPS (0-10) y las escalas no se recortan
    datos = _tabla(wb["Totales"], _bloques(wb["Totales"])["Opción"] + 1, 15, 2)
    assert len(datos) == 10 and all(d[1] is not None for d in datos)


def test_nps_grafica_sus_11_valores_sin_recorte(env):
    _, wb = _xlsx(env)
    ws = wb["Totales"]
    n_cats = []
    for ch in ws._charts:
        ref = (ch.series[0].cat.strRef or ch.series[0].cat.numRef).f.split("!")[1].replace("$", "")
        a, b = ref.split(":")
        n_cats.append(int(b[1:]) - int(a[1:]) + 1)
    assert 11 in n_cats and all(not t.startswith("10 opciones más frecuentes de 11") for t in _graficos(ws))


# ---------- pocos datos, bordes y errores ----------

def test_pocos_datos_en_grupo_de_menos_de_10_y_nota_en_resumen(env):
    _, wb = _xlsx(env, por="Tramo")
    ws = wb["Por grupo"]
    tabla = _tabla(ws, _bloques(ws)["Satisfacción"] + 2)
    assert {f[0]: f[2] for f in tabla} == {"Todas las personas": None, "20-24": "Sí", "25-29": "Sí", "30-34": None, "35-39": None}
    textos = [c.value for c in wb["Resumen"]["A"]]
    assert "Los grupos marcados con «Pocos datos» tienen menos de 10 personas: tómalos con cautela." in textos
    assert ["Separado por", "Tramo"] in [[c.value for c in f][:2] for f in wb["Resumen"].iter_rows()]
    # sin grupos chicos no hay nota
    _, wb = _xlsx(env, por="Género")
    assert not any(str(c.value).startswith("Los grupos marcados") for c in wb["Resumen"]["A"])


def test_resumen_con_unidos_y_tramos(env):
    unir = json.dumps([{"nombre": "20-29", "valores": ["20-24", "25-29"]}])
    _, wb = _xlsx(env, por="Tramo", unir=unir)
    textos = [c.value for c in wb["Resumen"]["A"]]
    assert "20-29: 20-24, 25-29" in textos and "Grupos unidos" in textos
    # con los dos grupos unidos ya no hay pocos datos (12 personas)
    assert not any(str(t).startswith("Los grupos marcados") for t in textos)
    _, wb = _xlsx(env, por="Edad", cortes="18,30,40")
    filas = [[c.value for c in f][:2] for f in wb["Resumen"].iter_rows()]
    assert ["Tramos usados", "18, 30, 40"] in filas and ["Personas sin grupo", 0] in filas


def test_errores_400_iguales_a_closed_results(env):
    for kw in (dict(por="no-existe"), dict(por="Género", unir="no-es-json"), dict(por="Edad", unir="[]"),
               dict(cortes="18,25"), dict(por="Edad", cortes="30,20"),
               dict(por="Tramo", unir=json.dumps([{"nombre": "x", "valores": ["20-24", "99"]}]))):
        r, _ = _xlsx(env, **kw)
        q = _query(por=env["ids"].get(kw.get("por"), kw.get("por")), cortes=kw.get("cortes"), unir=kw.get("unir"))
        esperado = env["c"].get(f"/api/projects/{env['pid']}/closed-results?{q}", headers=env["h"])
        assert r.status_code == 400 == esperado.status_code and r.json() == esperado.json(), kw
    assert _xlsx(env, por="no-existe")[0].json()["detail"] == "Esa pregunta no se puede usar para separar en grupos."


def test_lector_ok_otra_organizacion_404_y_sin_sesion_401(env):
    assert _xlsx(env, h="hl")[0].status_code == 200
    r, _ = _xlsx(env, h="hb")
    assert r.status_code == 404 and r.json()["detail"] == "Proyecto no encontrado"
    assert env["c"].get(f"/api/projects/{env['pid']}/export/closed-xlsx").status_code == 401


def test_nombre_del_archivo_y_tipo(env):
    r, _ = _xlsx(env)
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert r.headers["content-disposition"] == 'attachment; filename="resultados_cerradas_Encuesta_Nandu.xlsx"'
    r, _ = _xlsx(env, por="Género")
    assert r.headers["content-disposition"] == 'attachment; filename="resultados_cerradas_Encuesta_Nandu_por_Genero.xlsx"'


def test_proyecto_sin_cerradas_descarga_resumen_con_la_nota(env):
    c = env["c"]
    r = c.post("/api/projects", headers=env["h"], data={"name": "Solo texto", "preguntas": json.dumps(
        [{"titulo": "Comentario", "columnas": ["Comentario"], "tipo": "abierta"}])},
        files={"file": ("e.csv", io.BytesIO(("Comentario\n" + "\n".join(f"texto largo número {i}" for i in range(6))).encode()), "text/csv")})
    r = c.get(f"/api/projects/{r.json()['id']}/export/closed-xlsx", headers=env["h"])
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Resumen"] and export_closed.SIN_CERRADAS in [c.value for c in wb["Resumen"]["A"]]
