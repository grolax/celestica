"""T-017 RNF-16: los informes PDF/DOCX llevan la marca Celestica, no SoftInsight."""
from __future__ import annotations

import io
import os
import sys
from pathlib import Path

import mongomock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ.setdefault("ANTHROPIC_API_KEY", "sk-test")

from app import database  # noqa: E402

database._client = mongomock.MongoClient()

import docx  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from reportlab import rl_config  # noqa: E402

from app.main import app  # noqa: E402


def test_pdf_y_docx_dicen_celestica(monkeypatch):
    monkeypatch.setattr(rl_config, "pageCompression", 0)  # texto del PDF legible en bytes
    c = TestClient(app)
    r = c.post("/auth/signup", json={"email": "marca@x.cl", "password": "clave12345"})
    h = {"Authorization": f"Bearer {r.json()['token']}"}
    csv = "id_respuesta,respuesta\n" + "\n".join(f"E{i},Muy buen curso {i}" for i in range(1, 4))
    pid = c.post("/api/projects", headers=h, files={"file": ("d.csv", io.BytesIO(csv.encode()), "text/csv")},
                 data={"name": "Marca"}).json()["id"]

    pdf = c.get(f"/api/projects/{pid}/export/pdf", headers=h)
    assert pdf.status_code == 200
    assert pdf.headers["content-disposition"] == 'attachment; filename="informe_celestica_Marca.pdf"'
    assert b"Celestica" in pdf.content and b"SoftInsight" not in pdf.content

    dx = c.get(f"/api/projects/{pid}/export/docx", headers=h)
    assert dx.headers["content-disposition"] == 'attachment; filename="informe_celestica_Marca.docx"'
    d = docx.Document(io.BytesIO(dx.content))
    texto = "\n".join(p.text for p in d.paragraphs) + d.sections[0].footer.paragraphs[0].text
    assert "Celestica" in texto and "SoftInsight" not in texto
