"""
Exportación de reportes: PDF (reportlab), DOCX (python-docx), CSV (pandas).

El PDF/DOCX es el entregable que las universidades archivan para acreditación:
portada, resumen ejecutivo, análisis léxico, análisis temático con ejemplos,
tabla de trazabilidad completa y nota metodológica. Todo lo que aparece en el
reporte es auditable en la plataforma.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Optional

import pandas as pd

# ---------- Paleta Celestica (T-014) ----------
NAVY = "#0b2545"
PRIMARY = "#0077b6"
GRAY_ROW = "#F5F5F5"


# =====================================================================
# Helpers de datos (compartidos por PDF/DOCX/CSV)
# =====================================================================

def _cat_names(vista: dict) -> dict[str, str]:
    names = {c["id"]: c["nombre"] for c in vista["categories"]}
    names["SIN_CATEGORIA"] = "Sin categoría"
    return names


def _category_counts(vista: dict) -> list[tuple[str, int]]:
    names = _cat_names(vista)
    counts: dict[str, int] = {}
    for assignments in vista["codings"].values():
        for a in assignments:
            counts[a["category_id"]] = counts.get(a["category_id"], 0) + 1
    return sorted(((names.get(cid, cid), n) for cid, n in counts.items()), key=lambda x: -x[1])


def _confidence_note(conf) -> str:
    if not isinstance(conf, (int, float)):
        return ""
    if conf >= 0.85:
        return "Alta"
    if conf >= 0.60:
        return "Media"
    return "Baja"


def traceability_rows(vista: dict) -> list[dict]:
    names = _cat_names(vista)
    rows = []
    for r in vista["responses"]:
        assignments = vista["codings"].get(r["id"], [])
        if not assignments:
            rows.append({"id_respuesta": r["id"], "texto": r["text"], "categoria": "(no codificada)",
                         "justificacion": "", "confianza": None, "nivel_confianza": "",
                         "fuente": "", "estado": ""})
        for a in assignments:
            rows.append({
                "id_respuesta": r["id"], "texto": r["text"],
                "categoria": names.get(a["category_id"], a["category_id"]),
                "justificacion": a.get("rationale", ""),
                "confianza": a.get("confidence"),
                "nivel_confianza": _confidence_note(a.get("confidence")),
                "fuente": a.get("source", ""), "estado": a.get("status", ""),
            })
    return rows


def _category_examples(vista: dict, max_per_cat: int = 3) -> dict[str, list[dict]]:
    """Hasta N respuestas ejemplares por categoría (mayor confianza primero)."""
    by_cat: dict[str, list[dict]] = {}
    text_by_id = {r["id"]: r["text"] for r in vista["responses"]}
    for rid, assignments in vista["codings"].items():
        for a in assignments:
            by_cat.setdefault(a["category_id"], []).append({
                "text": text_by_id.get(rid, ""),
                "confidence": a.get("confidence") or 0,
                "rationale": a.get("rationale", ""),
            })
    return {
        cid: sorted(items, key=lambda x: -x["confidence"])[:max_per_cat]
        for cid, items in by_cat.items()
    }


def _freq_chart_png(vista: dict) -> Optional[bytes]:
    """Barras horizontales de los términos más frecuentes (matplotlib, backend Agg)."""
    lexical = vista.get("lexical")
    if not lexical or not lexical.get("top_terms"):
        return None
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return None
    terms = lexical["top_terms"][:15][::-1]
    fig, ax = plt.subplots(figsize=(7, 4.2))
    ax.barh([t["term"] for t in terms], [t["count"] for t in terms], color=PRIMARY)
    ax.set_title("Términos más frecuentes", fontsize=12)
    ax.tick_params(labelsize=9)
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    return buf.getvalue()


def _metadata(vista: dict) -> list[tuple[str, str]]:
    validated = [c for c in vista["categories"] if c["status"] == "validada"]
    return [
        ("Proyecto", vista["name"]),
        ("Fecha de generación", datetime.now().strftime("%d-%m-%Y %H:%M")),
        ("N° de respuestas analizadas", str(len(vista["responses"]))),
        ("N° de categorías validadas", str(len(validated))),
        ("N° de respuestas codificadas", str(len(vista["codings"]))),
        ("Marco teórico", vista.get("marco_teorico") or "(no especificado)"),
    ]


_METHODOLOGY = (
    "El análisis se realizó con Celestica en dos capas. Capa 1 (léxica): frecuencias, "
    "bigramas, co-ocurrencias y sentimiento por lexicón, calculados de forma determinística "
    "sin IA generativa. Capa 2 (temática): un modelo de lenguaje (Claude, Anthropic) propuso "
    "categorías ancladas al marco teórico del proyecto; el equipo analista validó, editó o "
    "descartó cada categoría antes de la codificación. La codificación asistida asignó cada "
    "respuesta a categorías validadas con justificación y nivel de confianza; el analista pudo "
    "re-codificar manualmente cualquier asignación. Toda decisión —humana o de IA— quedó "
    "registrada en el log de auditoría de la plataforma, disponible para revisión CNA/CNE. "
    "La tabla de trazabilidad de este documento vincula cada respuesta con su categoría, "
    "justificación, confianza y fuente (IA o analista)."
)


# =====================================================================
# CSV
# =====================================================================

def generate_csv_export(vista: dict) -> bytes:
    df = pd.DataFrame(traceability_rows(vista))
    buf = io.BytesIO()
    # utf-8-sig: BOM para que Excel en Windows abra bien acentos/ñ
    df.to_csv(buf, index=False, encoding="utf-8-sig")
    return buf.getvalue()


def generate_xlsx_export(vista: dict) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        pd.DataFrame(traceability_rows(vista)).to_excel(writer, sheet_name="Trazabilidad", index=False)
        pd.DataFrame(_category_counts(vista), columns=["Categoría", "N° respuestas"]).to_excel(
            writer, sheet_name="Resumen por categoría", index=False)
    return buf.getvalue()


# =====================================================================
# PDF (reportlab / platypus)
# =====================================================================

def generate_pdf_report(vista: dict, include_details: bool = True) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    primary = colors.HexColor(PRIMARY)
    navy = colors.HexColor(NAVY)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1x", parent=styles["Heading1"], textColor=navy, spaceAfter=10)
    h2 = ParagraphStyle("h2x", parent=styles["Heading2"], textColor=navy, spaceBefore=14)
    body = ParagraphStyle("bodyx", parent=styles["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("smallx", parent=body, fontSize=8, leading=10.5)
    cover_title = ParagraphStyle("coverx", parent=styles["Title"], fontSize=26, textColor=navy,
                                 alignment=TA_CENTER, spaceAfter=6)
    cover_sub = ParagraphStyle("coversubx", parent=styles["Normal"], fontSize=12, alignment=TA_CENTER,
                               textColor=colors.HexColor("#555555"))

    def _footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#888888"))
        canvas.drawString(2 * cm, 1.1 * cm, "Análisis realizado con Celestica · Auditable y trazable")
        canvas.drawRightString(letter[0] - 2 * cm, 1.1 * cm, f"Página {doc.page}")
        canvas.restoreState()

    def _table(data, col_widths=None, header=True, font_size=8):
        t = Table(data, colWidths=col_widths, repeatRows=1 if header else 0)
        style = [
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
            ("ROWBACKGROUNDS", (0, 1 if header else 0), (-1, -1),
             [colors.white, colors.HexColor(GRAY_ROW)]),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        if header:
            style += [
                ("BACKGROUND", (0, 0), (-1, 0), primary),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]
        t.setStyle(TableStyle(style))
        return t

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, topMargin=2 * cm, bottomMargin=2 * cm,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            title=f"Informe Celestica — {vista['name']}")
    story = []

    # ---- Portada ----
    story.append(Spacer(1, 5 * cm))
    story.append(Paragraph("Celestica", cover_title))
    story.append(Paragraph("Informe de Análisis Cualitativo", cover_sub))
    story.append(Spacer(1, 1.4 * cm))
    story.append(Paragraph(vista["name"], ParagraphStyle(
        "projname", parent=styles["Title"], fontSize=18, alignment=TA_CENTER)))
    story.append(Spacer(1, 2 * cm))
    meta_rows = [[Paragraph(f"<b>{k}</b>", small), Paragraph(str(v), small)] for k, v in _metadata(vista)]
    story.append(_table(meta_rows, col_widths=[5.5 * cm, 10.5 * cm], header=False, font_size=8.5))
    story.append(PageBreak())

    # ---- Resumen ejecutivo ----
    story.append(Paragraph("1. Resumen ejecutivo", h1))
    counts = _category_counts(vista)
    n_resp, n_coded = len(vista["responses"]), len(vista["codings"])
    resumen = (
        f"Se analizaron <b>{n_resp}</b> respuestas abiertas, de las cuales <b>{n_coded}</b> fueron "
        f"codificadas contra un libro de códigos de <b>{len([c for c in vista['categories'] if c['status'] == 'validada'])}</b> "
        f"categorías validadas por el equipo analista."
    )
    story.append(Paragraph(resumen, body))
    if counts:
        story.append(Paragraph("Distribución de respuestas por categoría:", body))
        story.append(Spacer(1, 6))
        story.append(_table([["Categoría", "N° respuestas"]] + [[c, str(n)] for c, n in counts],
                            col_widths=[11 * cm, 4 * cm]))

    # ---- Informe cualitativo (Fase 2) si existe ----
    if vista.get("qualitative_report"):
        story.append(Paragraph("2. Informe cualitativo", h1))
        for block in str(vista["qualitative_report"]).split("\n\n"):
            block = block.strip()
            if not block:
                continue
            if len(block) < 80 and block == block.upper():
                story.append(Paragraph(block, h2))
            else:
                story.append(Paragraph(block.replace("\n", "<br/>"), body))

    # ---- Análisis léxico ----
    story.append(Paragraph("3. Análisis léxico", h1))
    lexical = vista.get("lexical")
    if lexical:
        chart = _freq_chart_png(vista)
        if chart:
            story.append(Image(io.BytesIO(chart), width=15 * cm, height=9 * cm))
        top20 = lexical.get("top_terms", [])[:20]
        if top20:
            story.append(_table([["Término", "Frecuencia"]] + [[t["term"], str(t["count"])] for t in top20],
                                col_widths=[11 * cm, 4 * cm]))
        dist = lexical.get("sentiment_distribution") or {}
        if dist:
            story.append(Paragraph("Distribución de sentimiento:", body))
            story.append(_table([["Sentimiento", "N° respuestas"]] + [[k, str(v)] for k, v in dist.items()],
                                col_widths=[11 * cm, 4 * cm]))
        if vista.get("lexical_interpretation"):
            story.append(Paragraph("Interpretación léxica (IA):", h2))
            story.append(Paragraph(str(vista["lexical_interpretation"]), body))
    else:
        story.append(Paragraph("El análisis léxico no fue ejecutado en este proyecto.", body))

    # ---- Análisis temático ----
    story.append(Paragraph("4. Análisis temático — libro de códigos", h1))
    examples = _category_examples(vista)
    validated = [c for c in vista["categories"] if c["status"] == "validada"]
    for cat in validated:
        story.append(Paragraph(cat["nombre"], h2))
        if cat.get("definicion"):
            story.append(Paragraph(f"<b>Definición:</b> {cat['definicion']}", body))
        if cat.get("criterios_inclusion"):
            story.append(Paragraph(f"<b>Incluye:</b> {cat['criterios_inclusion']}", body))
        if cat.get("criterios_exclusion"):
            story.append(Paragraph(f"<b>Excluye:</b> {cat['criterios_exclusion']}", body))
        for ex in examples.get(cat["id"], []):
            story.append(Paragraph(f"«{ex['text']}» <i>(confianza {ex['confidence']:.2f})</i>", small))
    if not validated:
        story.append(Paragraph("No hay categorías validadas.", body))

    # ---- Trazabilidad ----
    if include_details:
        story.append(PageBreak())
        story.append(Paragraph("5. Tabla de trazabilidad", h1))
        story.append(Paragraph(
            "Cada fila vincula una respuesta con su categoría, justificación, confianza y fuente "
            "(IA o analista). Esta tabla es el insumo de auditoría CNA/CNE.", body))
        story.append(Spacer(1, 6))
        rows = traceability_rows(vista)
        header = ["ID", "Respuesta", "Categoría", "Justificación", "Conf.", "Fuente"]
        data = [header]
        for r in rows:
            conf = r["confianza"]
            data.append([
                Paragraph(str(r["id_respuesta"]), small),
                Paragraph(str(r["texto"])[:400], small),
                Paragraph(str(r["categoria"]), small),
                Paragraph(str(r["justificacion"])[:300], small),
                f"{conf:.2f}" if isinstance(conf, (int, float)) else "",
                r["fuente"],
            ])
        story.append(_table(data, col_widths=[1.6 * cm, 5.2 * cm, 3 * cm, 4.2 * cm, 1.2 * cm, 1.6 * cm], font_size=7))

    # ---- Metodología ----
    story.append(PageBreak())
    story.append(Paragraph("6. Nota metodológica", h1))
    story.append(Paragraph(_METHODOLOGY, body))

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


# =====================================================================
# DOCX (python-docx)
# =====================================================================

def generate_docx_report(vista: dict, include_details: bool = True) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    navy_rgb = RGBColor(0x0B, 0x25, 0x45)
    doc = docx.Document()

    def heading(text: str, level: int = 1):
        h = doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.color.rgb = navy_rgb
        return h

    def add_table(headers: list[str], rows: list[list[str]]):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for i, htext in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = htext
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)
                for p in cells[i].paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)
        return t

    # ---- Portada ----
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Celestica")
    run.font.size = Pt(34)
    run.font.color.rgb = navy_rgb
    run.bold = True
    sub = doc.add_paragraph("Informe de Análisis Cualitativo")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(vista["name"])
    name_run.font.size = Pt(20)
    name_run.bold = True
    doc.add_paragraph()
    add_table(["Campo", "Valor"], [[k, v] for k, v in _metadata(vista)])
    doc.add_page_break()

    # ---- Resumen ejecutivo ----
    heading("1. Resumen ejecutivo")
    counts = _category_counts(vista)
    doc.add_paragraph(
        f"Se analizaron {len(vista['responses'])} respuestas abiertas, de las cuales "
        f"{len(vista['codings'])} fueron codificadas contra un libro de códigos de "
        f"{len([c for c in vista['categories'] if c['status'] == 'validada'])} categorías validadas."
    )
    if counts:
        add_table(["Categoría", "N° respuestas"], [[c, str(n)] for c, n in counts])

    # ---- Informe cualitativo ----
    if vista.get("qualitative_report"):
        heading("2. Informe cualitativo")
        for block in str(vista["qualitative_report"]).split("\n\n"):
            block = block.strip()
            if not block:
                continue
            if len(block) < 80 and block == block.upper():
                heading(block.title(), level=2)
            else:
                doc.add_paragraph(block)

    # ---- Análisis léxico ----
    heading("3. Análisis léxico")
    lexical = vista.get("lexical")
    if lexical:
        chart = _freq_chart_png(vista)
        if chart:
            from docx.shared import Cm
            doc.add_picture(io.BytesIO(chart), width=Cm(15))
        top20 = lexical.get("top_terms", [])[:20]
        if top20:
            add_table(["Término", "Frecuencia"], [[t["term"], str(t["count"])] for t in top20])
        if vista.get("lexical_interpretation"):
            heading("Interpretación léxica (IA)", level=2)
            doc.add_paragraph(str(vista["lexical_interpretation"]))
    else:
        doc.add_paragraph("El análisis léxico no fue ejecutado en este proyecto.")

    # ---- Análisis temático ----
    heading("4. Análisis temático — libro de códigos")
    examples = _category_examples(vista)
    for cat in [c for c in vista["categories"] if c["status"] == "validada"]:
        heading(cat["nombre"], level=2)
        if cat.get("definicion"):
            doc.add_paragraph(f"Definición: {cat['definicion']}")
        if cat.get("criterios_inclusion"):
            doc.add_paragraph(f"Incluye: {cat['criterios_inclusion']}")
        if cat.get("criterios_exclusion"):
            doc.add_paragraph(f"Excluye: {cat['criterios_exclusion']}")
        for ex in examples.get(cat["id"], []):
            p = doc.add_paragraph(style="Intense Quote")
            p.add_run(f"«{ex['text']}» (confianza {ex['confidence']:.2f})").font.size = Pt(9)

    # ---- Trazabilidad ----
    if include_details:
        doc.add_page_break()
        heading("5. Tabla de trazabilidad")
        doc.add_paragraph(
            "Cada fila vincula una respuesta con su categoría, justificación, confianza y fuente. "
            "Insumo de auditoría CNA/CNE."
        )
        rows = traceability_rows(vista)
        add_table(
            ["ID", "Respuesta", "Categoría", "Justificación", "Conf.", "Fuente"],
            [[r["id_respuesta"], str(r["texto"])[:400], r["categoria"], str(r["justificacion"])[:300],
              f"{r['confianza']:.2f}" if isinstance(r["confianza"], (int, float)) else "", r["fuente"]]
             for r in rows],
        )

    # ---- Metodología ----
    doc.add_page_break()
    heading("6. Nota metodológica")
    doc.add_paragraph(_METHODOLOGY)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Análisis realizado con Celestica · Auditable y trazable"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
