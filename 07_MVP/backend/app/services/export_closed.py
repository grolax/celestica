"""RF-30 / RF-31 (T-027): Excel con los resultados de las preguntas cerradas, con gráficos nativos.

No calcula nada: ordena y grafica el resultado de `quantitative.calcular` (las mismas cifras de `closed-results`).
Las reglas están en docs/API_CONTRATOS.md.
"""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, PatternFill

from app.services.export import NAVY, PRIMARY
from app.services.quantitative import CERRADAS, CORTES_DEFECTO, orden_natural

NOTA = ("Cifras calculadas por la plataforma directamente desde las respuestas, sin IA. Los porcentajes son sobre "
        "quienes respondieron cada pregunta; quienes no respondieron aparecen como N/A.")
SIN_CERRADAS = "Este proyecto no tiene preguntas cerradas."
DEMASIADOS = "Hay demasiados grupos para un gráfico legible; revisa la tabla."
TIPOS = {"escala": "Escala", "nps": "NPS", "si_no": "Sí / No", "seleccion_unica": "Selección única",
         "seleccion_multiple": "Selección múltiple", "segmento": "Segmento", "numerica": "Numérica"}
# Tokens de marca de frontend/src/styles/globals.css (navy, primary, warning, success, celeste, danger).
COLORES = [PRIMARY, NAVY, "#b45309", "#0f7a2f", "#0099cc", "#b91c1c"]
MAX_OPCIONES, MAX_SERIES, MAX_GRUPOS = 10, 6, 15
ALTO_GRAFICO = 17  # filas que ocupa un gráfico de 8 cm


def _fila(ws, row: int, vals, col: int = 1, cabecera: bool = False) -> None:
    for i, v in enumerate(vals):
        c = ws.cell(row, col + i, v)
        if cabecera:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY[1:])


def _titulo(ws, row: int, texto: str) -> None:
    ws.cell(row, 1, texto).font = Font(bold=True, size=12, color=NAVY[1:])


def _grafico(titulo, cats, series, horizontal=False, linea=False):
    """`series` = [(Reference de valores, nombre)]; siempre con el valor escrito en cada barra o punto."""
    ch = LineChart() if linea else BarChart()
    if not linea:
        ch.type, ch.grouping = ("bar" if horizontal else "col"), "clustered"
    ch.title, ch.width, ch.height = titulo, 18, 8
    ch.x_axis.delete = ch.y_axis.delete = False  # openpyxl 3.1 los oculta por defecto
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"  # la primera opción arriba
    for i, (vals, nombre) in enumerate(series):
        s = Series(vals, title=nombre)
        color = COLORES[i % len(COLORES)][1:]
        if linea:
            s.graphicalProperties.line.solidFill, s.marker.symbol, s.smooth = color, "circle", False
            s.marker.graphicalProperties.solidFill = color
        else:
            s.graphicalProperties.solidFill = color
        ch.series.append(s)
    ch.set_categories(cats)
    ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False,
                                  showPercent=False, dLblPos="t" if linea else "outEnd")
    if len(series) == 1:
        ch.legend = None
    return ch


def _extras(p: dict) -> list[tuple]:
    """Filas de estadísticos de una pregunta según su tipo (etiqueta, valor[, %])."""
    if e := p.get("escala"):
        if p["tipo"] == "nps":
            n = p["nps"]
            return [("NPS", n["valor"]), *[(nom, n[k]["n"], n[k]["pct"]) for nom, k in
                                          (("Promotores", "promotores"), ("Pasivos", "pasivos"), ("Detractores", "detractores"))]]
        t, b = e["top2"], e["bottom2"]
        return [("Promedio", e["promedio"]), ("Desviación", e["desviacion"]),
                (f"Notas altas ({t['valores'][0]} y {t['valores'][1]})", t["n"], t["pct"]),
                (f"Notas bajas ({b['valores'][0]} y {b['valores'][1]})", b["n"], b["pct"])]
    if n := p.get("numerica"):
        return [("Promedio", n["promedio"]), ("Mediana", n["mediana"]), ("Mínimo", n["min"]), ("Máximo", n["max"])]
    return [(p["nota"],)] if p.get("nota") else []


def _bloque_totales(ws, p: dict, row: int) -> int:
    """Escribe el bloque de una pregunta y devuelve la fila donde empieza el siguiente."""
    ini = row
    _titulo(ws, row, p["titulo"])
    for i, (k, v) in enumerate([("Tipo", TIPOS[p["tipo"]]), ("Respondieron", p["n_respuestas"]), ("N/A", p["n_na"]),
                                ("% de respuesta", p["tasa_respuesta"])], 1):
        _fila(ws, row + i, (k, v))
    row += 6
    ops = p.get("opciones", [])
    if ops:
        _fila(ws, row, ("Opción", "Personas", "%"), cabecera=True)
        for i, o in enumerate(ops, 1):
            _fila(ws, row + i, (o["valor"], o["n"], o["pct"]))
        ref, fin = row + 1, row + len(ops)
        if len(ops) > MAX_OPCIONES:  # datos del gráfico aparte: las 10 más frecuentes, en el orden de la tabla
            top = sorted(sorted(range(len(ops)), key=lambda i: -ops[i]["n"])[:MAX_OPCIONES])
            _fila(ws, ini, ("Datos del gráfico", "%"), col=15, cabecera=True)
            for j, i in enumerate(top, 1):
                _fila(ws, ini + j, (ops[i]["valor"], ops[i]["pct"]), col=15)
            cats, vals = Reference(ws, min_col=15, min_row=ini + 1, max_row=ini + MAX_OPCIONES), \
                Reference(ws, min_col=16, min_row=ini + 1, max_row=ini + MAX_OPCIONES)
            titulo = f"{MAX_OPCIONES} opciones más frecuentes de {len(ops)}"
        else:
            cats, vals, titulo = Reference(ws, min_col=1, min_row=ref, max_row=fin), \
                Reference(ws, min_col=3, min_row=ref, max_row=fin), "Porcentaje por opción"
        ws.add_chart(_grafico(titulo, cats, [(vals, "%")], horizontal=True), f"E{ini + 1}")
        row = fin + 2
    for fila in _extras(p):
        _fila(ws, row, fila)
        row += 1
    return max(row, ini + ALTO_GRAFICO) + 2


def _columnas(p: dict) -> list[tuple]:
    """(encabezado, función que saca el valor de una fila total o de grupo) según el tipo, como la pantalla."""
    g = lambda *ks: lambda i: _dig(i, ks)  # noqa: E731
    if p["tipo"] == "nps":
        return [("NPS", g("nps", "valor")), ("Promotores (%)", g("nps", "promotores", "pct")),
                ("Detractores (%)", g("nps", "detractores", "pct"))]
    if p["tipo"] == "escala":
        return [("Promedio", g("escala", "promedio")), ("Notas altas (%)", g("escala", "top2", "pct")),
                ("Notas bajas (%)", g("escala", "bottom2", "pct"))]
    if p["tipo"] == "numerica":
        return [(h, g("numerica", k)) for h, k in (("Promedio", "promedio"), ("Mediana", "mediana"),
                                                   ("Mínimo", "min"), ("Máximo", "max"))]
    return [(f"{o['valor']} (%)", lambda i, k=k: i["opciones"][k]["pct"]) for k, o in enumerate(p["opciones"])]


def _dig(d, ks):
    for k in ks:
        d = d.get(k) if d else None
    return d


def _bloque_grupos(ws, p: dict, res: dict, natural: bool, row: int) -> int:
    ini = row
    _titulo(ws, row, p["titulo"])
    cols = _columnas(p)
    _fila(ws, row + 1, ["Grupo", "Respondieron", "Pocos datos"] + [h for h, _ in cols], cabecera=True)
    filas = [("Todas las personas", p, False)] + [(g["valor"], pg, g["pocos_datos"])
                                                  for g, pg in zip(res["grupos"], p["por_grupo"])]
    for i, (nombre, item, poco) in enumerate(filas, 2):
        _fila(ws, row + i, [nombre, item["n_respuestas"], "Sí" if poco else None] + [f(item) for _, f in cols])
    primero, ultimo = row + 3, row + 1 + len(filas)  # sin la fila "Todas las personas"
    row = ultimo + 2
    if len(res["grupos"]) > MAX_GRUPOS:
        ws.cell(row, 1, DEMASIADOS)
        return row + 2
    cats = Reference(ws, min_col=1, min_row=primero, max_row=ultimo)
    col = lambda c: Reference(ws, min_col=c, min_row=primero, max_row=ultimo)  # noqa: E731
    if p["tipo"] in ("escala", "nps", "numerica"):
        nom = cols[0][0]
        graf = _grafico(f"{nom} por grupo" + (" (tendencia)" if natural else ""), cats, [(col(4), nom)], linea=natural)
    else:
        n = min(len(cols), MAX_SERIES)
        graf = _grafico("Porcentaje de cada opción por grupo" + (f" ({n} primeras opciones)" if len(cols) > n else ""),
                        cats, [(col(4 + k), p["opciones"][k]["valor"]) for k in range(n)])
    ws.add_chart(graf, f"A{row}")
    return row + ALTO_GRAFICO + 2


def _resumen(ws, project: dict, res: dict, cortes: str | None, sin_cerradas: bool) -> None:
    _titulo(ws, 1, "Resultados de preguntas cerradas")
    filas = [("Proyecto", project["name"]), ("Fecha de descarga", datetime.now().strftime("%d-%m-%Y")),
             ("Personas", res["n_personas"]), (NOTA,)]
    if sin_cerradas:
        filas.append((SIN_CERRADAS,))
    for i, f in enumerate(filas, 3):
        _fila(ws, i, f)
    if "por" not in res:
        return
    row = 3 + len(filas) + 1
    _fila(ws, row, ("Separado por", res["por"]["titulo"]))
    if res["por"]["tipo"] == "numerica":
        _fila(ws, row + 1, ("Tramos usados", (cortes if cortes is not None else CORTES_DEFECTO).replace(",", ", ")))
        row += 1
    _fila(ws, row + 1, ("Personas sin grupo", res["n_sin_grupo"]))
    _fila(ws, row + 3, ("Grupo", "Personas"), cabecera=True)
    row += 4
    for g in res["grupos"]:
        _fila(ws, row, (g["valor"], g["n_personas"]))
        row += 1
    unidos = [f"{g['valor']}: {', '.join(g['unidos'])}" for g in res["grupos"] if "unidos" in g]
    if unidos:
        _fila(ws, row + 1, ("Grupos unidos",))
        for i, u in enumerate(unidos, 2):
            _fila(ws, row + i, (u,))
        row += len(unidos) + 2
    if any(g["pocos_datos"] for g in res["grupos"]):
        _fila(ws, row + 1, (f"Los grupos marcados con «Pocos datos» tienen menos de {res['umbral_pocos_datos']} "
                            "personas: tómalos con cautela.",))


def generate_closed_xlsx(project: dict, res: dict, cortes: str | None = None) -> bytes:
    """`res` = `quantitative.calcular(...)` del proyecto con los mismos parámetros que la descarga."""
    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"
    resumen.column_dimensions["A"].width = 30
    resumen.column_dimensions["B"].width = 40
    _resumen(resumen, project, res, cortes, not any(q["tipo"] in CERRADAS for q in project.get("preguntas", [])))
    if res["preguntas"]:
        ws = wb.create_sheet("Totales")
        ws.column_dimensions["A"].width = 44
        row = 1
        for p in res["preguntas"]:
            row = _bloque_totales(ws, p, row)
        if "por" in res:
            ws = wb.create_sheet("Por grupo")
            ws.column_dimensions["A"].width = 30
            natural, row = orden_natural(res), 1
            for p in res["preguntas"]:
                row = _bloque_grupos(ws, p, res, natural, row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
